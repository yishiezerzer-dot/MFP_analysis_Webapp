"""End-to-end: the one-click LCMS pipeline as a background job over HTTP, on the example data."""
import io
import json
import threading
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.automation.registry import list_actions
from app.db import get_job
from app.main import app
from app.services.jobs import JobManager

client = TestClient(app)
DATA = Path(__file__).resolve().parents[1] / "app" / "examples" / "data"
SETTINGS = {
    "monomers": [{"name": "GA", "mass": 76.016044}, {"name": "LA", "mass": 90.031694}],
    "polarity": "positive",
    "tolerance_da": 0.5,
    "adducts": ["H", "Na", "Cl"],
    "rules": {"elution_direction": "increasing"},
}


def wait_for(job_id: str, timeout: float = 120.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        status = client.get(f"/api/lcms/pipeline/{job_id}").json()
        if status["status"] in ("done", "failed", "cancelled"):
            return status
        time.sleep(0.2)
    raise AssertionError(f"job {job_id} did not finish: {status}")


def start(settings=SETTINGS) -> str:
    files = {
        "mzml": ("example_plga.mzML", (DATA / "example_plga.mzML").read_bytes(), "application/octet-stream"),
        "uv": ("example_plga_uv.csv", (DATA / "example_plga_uv.csv").read_bytes(), "text/csv"),
    }
    resp = client.post("/api/lcms/pipeline", data={"settings": json.dumps(settings)}, files=files, headers={"X-Workspace-Id": "pipe_ws"})
    assert resp.status_code == 200, resp.text
    return resp.json()["job_id"]


def start_on_session(session_id: str, settings=SETTINGS) -> str:
    resp = client.post("/api/lcms/pipeline", data={"settings": json.dumps(settings), "session_id": session_id},
                       headers={"X-Workspace-Id": "pipe_ws"})
    assert resp.status_code == 200, resp.text
    return resp.json()["job_id"]


@pytest.fixture(scope="module")
def finished_job() -> str:
    job_id = start()
    status = wait_for(job_id)
    assert status["status"] == "done", status
    return job_id


def test_job_runs_to_the_end_and_reports_progress(finished_job):
    status = client.get(f"/api/lcms/pipeline/{finished_job}").json()
    assert status["progress"] == 1.0 and status["step"] == "Done"
    assert any("UV→MS offset" in m for m in status["messages"])
    assert status["session_id"]
    # The run is an ordinary LCMS session with its UV attached.
    session = client.get(f"/api/lcms/sessions/{status['session_id']}").json()
    assert session["uv"]["available"] is True


def test_result_has_aligned_peaks_and_checked_labels(finished_job):
    result = client.get(f"/api/lcms/pipeline/{finished_job}/result").json()
    assert result["alignment"]["offset_min"] == pytest.approx(0.12, abs=0.02)
    assert len(result["peaks"]) >= 20
    approved = [lab for lab in result["labels"] if lab["review"] == "approved"]
    assert any(lab["species"] == "1-GA + 1-LA" for lab in approved)
    for lab in approved:
        assert lab["status"] == "ok" and lab["layout"]["axRef"] == "pixel"
    assert {lab["status"] for lab in result["labels"]} >= {"ok", "flagged", "alternative"}
    assert all(isinstance(lab["reasons"], list) for lab in result["labels"])


def test_figures_and_workbook_download(finished_job):
    svg = client.get(f"/api/lcms/pipeline/{finished_job}/figure.svg")
    assert svg.status_code == 200 and svg.headers["content-type"].startswith("image/svg+xml") and b"<svg" in svg.content
    png = client.get(f"/api/lcms/pipeline/{finished_job}/figure.png")
    assert png.status_code == 200 and png.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert b"<g id=\"grid" not in svg.content  # no gridlines
    # The print version: 180 × 100 mm at 600 dpi (IHDR width and height in pixels).
    print_png = client.get(f"/api/lcms/pipeline/{finished_job}/figure-print.png")
    assert print_png.status_code == 200
    width, height = int.from_bytes(print_png.content[16:20], "big"), int.from_bytes(print_png.content[20:24], "big")
    assert abs(width - round(180 / 25.4 * 600)) <= 2 and abs(height - round(100 / 25.4 * 600)) <= 2
    xlsx = client.get(f"/api/lcms/pipeline/{finished_job}/peaks.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Peaks", "Candidates", "Settings"]
    settings = {row[0]: row[1] for row in wb["Settings"].iter_rows(min_row=2, values_only=True)}
    assert settings["tolerance_da"] == 0.5 and settings["provenance.mzml_sha256"]
    assert "provenance.app_version" in settings


def test_review_regenerates_outputs(finished_job):
    result = client.get(f"/api/lcms/pipeline/{finished_job}/result").json()
    approved = next(lab for lab in result["labels"] if lab["review"] == "approved")
    pending = next(lab for lab in result["labels"] if lab["review"] == "pending")
    before = client.get(f"/api/lcms/pipeline/{finished_job}/figure.svg").content
    resp = client.patch(
        f"/api/lcms/pipeline/{finished_job}/labels",
        json={"updates": [{"id": approved["id"], "review": "rejected"}, {"id": pending["id"], "review": "approved"}]},
    )
    assert resp.status_code == 200, resp.text
    labels = {lab["id"]: lab for lab in resp.json()["labels"]}
    assert labels[approved["id"]]["review"] == "rejected" and labels[approved["id"]]["layout"] is None
    assert labels[pending["id"]]["review"] == "approved" and labels[pending["id"]]["layout"] is not None
    assert client.get(f"/api/lcms/pipeline/{finished_job}/figure.svg").content != before
    bad = client.patch(f"/api/lcms/pipeline/{finished_job}/labels", json={"updates": [{"id": "nope", "review": "approved"}]})
    assert bad.status_code == 422


def test_peak_spectrum_is_the_window_the_label_came_from(finished_job):
    result = client.get(f"/api/lcms/pipeline/{finished_job}/result").json()
    label = next(lab for lab in result["labels"] if lab["species"] == "1-GA + 1-LA" and lab["rank"] == 1)
    peak = next(p for p in result["peaks"] if p["id"] == label["peak_id"])
    sp = client.get(f"/api/lcms/pipeline/{finished_job}/peaks/{peak['id']}/spectrum").json()
    assert sp["meta"]["spectrum_id"] == f"pipeline:{peak['id']}"
    assert sp["meta"]["rt_start"] < peak["ms_rt_min"] < sp["meta"]["rt_end"]
    assert sp["meta"]["merge_mode"].startswith("pipeline") and sp["meta"]["n_scans"] >= 1
    assert len(sp["mz"]) == len(sp["intensity"]) > 0
    assert any(lab["text"] == label["text"] and abs(lab["mz"] - label["mz"]) < 1e-6 for lab in sp["labels"])
    assert client.get(f"/api/lcms/pipeline/{finished_job}/peaks/P999/spectrum").status_code == 404


def test_bad_requests_are_explained():
    resp = client.post("/api/lcms/pipeline", data={"settings": json.dumps({"monomers": []}), "session_id": "x"})
    assert resp.status_code == 422 and "monomers" in resp.json()["detail"]
    resp = client.post("/api/lcms/pipeline", data={"settings": "{not json"})
    assert resp.status_code == 422 and "JSON" in resp.json()["detail"]
    resp = client.post("/api/lcms/pipeline", data={"settings": json.dumps(SETTINGS)})
    assert resp.status_code == 422 and "mzML" in resp.json()["detail"]
    assert client.get("/api/lcms/pipeline/missing").status_code == 404


def test_cancelling_stops_a_running_job():
    manager = JobManager()
    started = threading.Event()

    def slow(ctx):
        started.set()
        for i in range(200):
            ctx.progress(i / 200, f"step {i}")
            time.sleep(0.02)

    job_id = manager.submit("pipeline", slow, params={}, workspace_id="pipe_ws")
    assert started.wait(5)
    assert manager.cancel(job_id) is True
    deadline = time.time() + 5
    while get_job(job_id)["status"] == "running" and time.time() < deadline:
        time.sleep(0.05)
    assert get_job(job_id)["status"] == "cancelled"
    assert manager.cancel(job_id) is False  # already finished


def test_automation_action_is_registered():
    assert "lcms.run_pipeline" in {spec.id for spec in list_actions()}


def test_past_runs_are_listed_and_deleted_with_the_file_they_opened():
    job_id = start()
    assert wait_for(job_id)["status"] == "done"
    runs = client.get("/api/lcms/pipeline", headers={"X-Workspace-Id": "pipe_ws"}).json()
    run = next(r for r in runs if r["job_id"] == job_id)
    assert run["status"] == "done" and run["uploaded"] is True and run["file_open"] is True
    assert run["display_name"] == "example_plga.mzML" and run["n_peaks"] >= 20 and run["approved"] >= 1
    out_dir = Path(get_job(job_id)["output_dir"])
    sid = run["session_id"]
    deleted = client.delete(f"/api/lcms/pipeline/{job_id}/run").json()
    assert deleted == {"deleted": True, "file_closed": True}
    assert not out_dir.exists() and get_job(job_id) is None
    assert client.get(f"/api/lcms/sessions/{sid}").status_code == 404  # the run's own upload is closed
    assert all(r["job_id"] != job_id for r in client.get("/api/lcms/pipeline", headers={"X-Workspace-Id": "pipe_ws"}).json())
    assert client.delete(f"/api/lcms/pipeline/{job_id}/run").status_code == 404


def test_a_run_on_an_already_open_file_leaves_the_file_open(finished_job):
    sid = client.get(f"/api/lcms/pipeline/{finished_job}").json()["session_id"]
    job_id = start_on_session(sid)
    assert wait_for(job_id)["status"] == "done"
    assert client.delete(f"/api/lcms/pipeline/{job_id}/run").json() == {"deleted": True, "file_closed": False}
    assert client.get(f"/api/lcms/sessions/{sid}").status_code == 200


def test_runs_older_than_five_days_are_cleaned_up():
    from datetime import datetime, timedelta, timezone

    from app.db import get_db_connection
    from app.routers.pipeline import cleanup_old_runs

    old, recent = start(), start()
    assert wait_for(old)["status"] == "done" and wait_for(recent)["status"] == "done"
    stamp = (datetime.now(timezone.utc) - timedelta(days=6)).isoformat()
    conn = get_db_connection()
    with conn:
        conn.execute("UPDATE jobs SET created_at = ? WHERE id = ?", (stamp, old))
    conn.close()
    assert cleanup_old_runs(days=5) >= 1
    assert get_job(old) is None and get_job(recent) is not None


def test_a_running_job_cannot_be_deleted():
    manager = JobManager()
    started = threading.Event()

    def slow(ctx):
        started.set()
        for i in range(200):
            ctx.progress(i / 200, f"step {i}")
            time.sleep(0.02)

    job_id = manager.submit("pipeline", slow, params={}, workspace_id="pipe_ws")
    assert started.wait(5)
    assert client.delete(f"/api/lcms/pipeline/{job_id}/run").status_code == 409
    manager.cancel(job_id)

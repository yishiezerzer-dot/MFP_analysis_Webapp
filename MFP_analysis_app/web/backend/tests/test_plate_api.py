"""Plate Reader API: Gen5 upload, layouts, analysis, templates, experiments, calculation workbook."""
import io
import re
import statistics as st
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)
FIX = Path(__file__).parent / "fixtures" / "plate_reader"
WS = {"X-Workspace-Id": "plate_api_ws"}


def upload(name):
    data = (FIX / name).read_bytes()
    resp = client.post("/api/plate-reader/sessions", files={"file": (name, data, "application/octet-stream")}, headers=WS)
    assert resp.status_code == 200, resp.text
    return resp.json()


def gent_layout():
    return {
        "dilution": {"top": 1024, "unit": "µg/mL", "factor": 2, "direction": "columns", "first": 1, "last": 11},
        "groups": [{"id": "g1", "name": "Gentamicin", "kind": "reference", "wells": [f"{r}{c}" for r in "ABC" for c in range(1, 12)]}],
        "growth_control": ["A12", "B12", "C12"],
        "blank": [f"{r}{c}" for r in "DEFGH" for c in range(1, 13)],
        "excluded": ["B12"],
    }


def test_upload_returns_plate_values_metadata_and_layout_from_notes():
    s = upload("gen5_lacglydoh_511_111.xlsx")
    assert s["metadata"]["plate_number"] == "Plate 1"
    assert s["values"]["F1"] == pytest.approx(0.485)
    assert s["notes"] == ["A-C - LacGlyDOH 5:1:1", "D-F - LacGlyDOH 1:1:1"]
    assert s["layout_source"] == "notes"
    assert [g["name"] for g in s["layout"]["groups"]] == ["LacGlyDOH 5:1:1", "LacGlyDOH 1:1:1"]
    assert s["layout"]["growth_control"] == [f"{r}12" for r in "ABCDEF"]


def test_non_plate_file_is_rejected_with_a_clear_message():
    resp = client.post("/api/plate-reader/sessions", files={"file": ("x.csv", b"a,b\n1,2\n", "text/csv")}, headers=WS)
    assert resp.status_code == 400
    assert "8×12" in resp.json()["detail"]


def test_layout_is_saved_and_analysis_is_recorded():
    sid = upload("gen5_gentamicin.xlsx")["session_id"]
    put = client.put(f"/api/plate-reader/sessions/{sid}/layout", json=gent_layout())
    assert put.status_code == 200
    s = client.get(f"/api/plate-reader/sessions/{sid}").json()
    assert s["layout_source"] == "saved" and s["layout"]["excluded"] == ["B12"]

    res = client.post(f"/api/plate-reader/sessions/{sid}/analysis", json={"subtract_blank": True}).json()
    pct = [round(p["percent_growth"]) for p in res["groups"][0]["points"]]
    assert pct == [0, 0, 0, 0, 0, 0, 0, 0, 0, 22, 46]
    recorded = client.get("/api/experiments/results", params={"session_id": sid}).json()
    assert any(r["kind"] == "mic_plate" for r in recorded)


def test_invalid_layout_is_rejected():
    sid = upload("gen5_gentamicin.xlsx")["session_id"]
    bad = gent_layout()
    bad["excluded"] = ["Z99"]
    assert client.put(f"/api/plate-reader/sessions/{sid}/layout", json=bad).status_code == 422


def test_templates_crud():
    body = {"name": "MIC gentamicin A–C", "layout": gent_layout()}
    created = client.post("/api/plate-reader/templates", json=body).json()
    assert created["name"] == body["name"]
    body["layout"]["excluded"] = []
    again = client.post("/api/plate-reader/templates", json=body).json()
    assert again["id"] == created["id"] and again["layout"]["excluded"] == []
    names = [t["name"] for t in client.get("/api/plate-reader/templates").json()]
    assert body["name"] in names
    assert client.delete(f"/api/plate-reader/templates/{created['id']}").status_code == 200
    assert body["name"] not in [t["name"] for t in client.get("/api/plate-reader/templates").json()]


def test_experiment_combines_plates_with_the_same_tag():
    tag = "MIC run test"
    poly = upload("gen5_lacglydoh_511_111.xlsx")["session_id"]
    gent = upload("gen5_gentamicin.xlsx")["session_id"]
    client.put(f"/api/plate-reader/sessions/{gent}/layout", json=gent_layout())
    for sid in (poly, gent):
        client.put(f"/api/experiments/sessions/{sid}/tag", json={"experiment_tag": tag})
    exp = client.get(f"/api/plate-reader/experiments/{tag}").json()
    names = [g["name"] for p in exp["plates"] for g in p["analysis"]["groups"]]
    assert names == ["LacGlyDOH 5:1:1", "LacGlyDOH 1:1:1", "Gentamicin"] or set(names) == {"LacGlyDOH 5:1:1", "LacGlyDOH 1:1:1", "Gentamicin"}


def test_old_wizard_endpoints_are_gone():
    sid = upload("gen5_gentamicin.xlsx")["session_id"]
    assert client.post(f"/api/plate-reader/sessions/{sid}/mic", json={}).status_code in (404, 405)
    assert client.post(f"/api/plate-reader/sessions/{sid}/load", json={}).status_code in (404, 405)


# --- calculation workbook: formulas must recompute to the API numbers --------------------

_REF = re.compile(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)|\$?\b([A-Z]{1,2})\$?(\d+)\b")


def _evaluate(wb, sheet, formula, depth=0):
    """Tiny evaluator for the functions the workbook uses: AVERAGE, STDEV.S, COUNT, + - * /."""
    assert depth < 20
    expr = formula.lstrip("=")

    def cell(sh, col, row):
        v = wb[sh][f"{col}{row}"].value
        if isinstance(v, str) and v.startswith("="):
            return _evaluate(wb, sh, v, depth + 1)
        return v

    def args(text):
        vals = []
        for m in _REF.finditer(text):
            if m.group(1):
                vals.append(cell(m.group(1), m.group(2), m.group(3)))
            else:
                vals.append(cell(sheet, m.group(4), m.group(5)))
        return [v for v in vals if isinstance(v, (int, float))]

    def fn(match):
        name, inner = match.group(1), match.group(2)
        vals = args(inner)
        if name == "AVERAGE":
            return repr(st.mean(vals))
        if name == "STDEV.S":
            return repr(st.stdev(vals))
        if name == "COUNT":
            return repr(len(vals))
        raise AssertionError(name)

    while re.search(r"(AVERAGE|STDEV\.S|COUNT)\(([^()]*)\)", expr):
        expr = re.sub(r"(AVERAGE|STDEV\.S|COUNT)\(([^()]*)\)", fn, expr)
    expr = _REF.sub(lambda m: repr(cell(m.group(1), m.group(2), m.group(3)) if m.group(1) else cell(sheet, m.group(4), m.group(5))), expr)
    return eval(expr, {"__builtins__": {}}, {})  # noqa: S307 - test-only arithmetic on numbers


def test_calculation_workbook_recomputes_to_the_api_values():
    sid = upload("gen5_gentamicin.xlsx")["session_id"]
    client.put(f"/api/plate-reader/sessions/{sid}/layout", json=gent_layout())
    api = client.post(f"/api/plate-reader/sessions/{sid}/analysis", json={"subtract_blank": True}).json()
    resp = client.get(f"/api/plate-reader/sessions/{sid}/workbook", params={"subtract_blank": True})
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert {"Raw plate", "Layout", "Calculation", "Checks"} <= set(wb.sheetnames)
    calc = wb["Calculation"]
    header = [c.value for c in calc[1]]
    rows = [r for r in calc.iter_rows(min_row=2) if r[0].value == "Gentamicin"]
    col = {name: i for i, name in enumerate(header)}
    api_points = api["groups"][0]["points"]
    assert len(rows) == len(api_points)
    for row, point in zip(rows, api_points):
        assert row[col["Concentration"]].value == pytest.approx(point["concentration"])
        mean = _evaluate(wb, "Calculation", row[col["Mean OD (blank-subtracted)"]].value)
        pct = _evaluate(wb, "Calculation", row[col["% growth"]].value)
        assert mean == pytest.approx(point["mean"], abs=1e-9)
        assert pct == pytest.approx(point["percent_growth"], abs=1e-9)

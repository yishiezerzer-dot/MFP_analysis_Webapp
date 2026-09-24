"""The SI package must contain only recorded results, identical to what the app computed."""
import hashlib
import io
import json
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.main import app
from test_mzml_fast_index import _generate_synthetic_mzml

client = TestClient(app)
TAG = "SI_real_results"
CONC = [64 / 2**i for i in range(11)] + [0.0]


def _tag(sid):
    client.put(f"/api/experiments/sessions/{sid}/tag", json={"experiment_tag": TAG})


@pytest.fixture(scope="module")
def run():
    x = np.linspace(600, 4000, 1200)
    ftir_bytes = "\n".join(
        f"{u:.2f},{v:.6f}" for u, v in zip(x, 0.8 * np.exp(-0.5 * ((x - 1735) / 10) ** 2) + 0.3 * np.exp(-0.5 * ((x - 1650) / 12) ** 2) + 0.05)
    ).encode()
    ftir = client.post("/api/ftir/sessions", files={"file": ("pla_film.csv", ftir_bytes, "text/csv")}).json()["session_id"]
    _tag(ftir)
    peaks = client.post(
        f"/api/ftir/sessions/{ftir}/peaks",
        json={"mode": "absorbance", "baseline": "airpls", "baseline_lambda": 12345.0, "min_prominence": 0.05, "assign": True},
    ).json()
    fit = client.post(
        f"/api/ftir/sessions/{ftir}/fit",
        json={"mode": "absorbance", "region": [1600, 1780], "n_components": 2, "profile": "voigt"},
    ).json()

    od = lambda c: 1.0 if c == 0 else 0.05 + 0.95 / (1 + (c / 2.0) ** 2)
    rng = np.random.default_rng(1)
    plate_df = pd.DataFrame([[od(c) + rng.normal(0, 0.01) for c in CONC] for _ in range(3)], columns=[str(i) for i in range(1, 13)])
    plate_bytes = plate_df.to_csv(index=False).encode()
    plate = client.post("/api/plate-reader/sessions", files={"file": ("mic_plate.csv", plate_bytes, "text/csv")}).json()["session_id"]
    _tag(plate)
    mic = client.post(
        f"/api/plate-reader/sessions/{plate}/mic",
        json={"sample_rows": [0, 1, 2], "concentration_columns": [str(i) for i in range(1, 13)],
              "tick_text": ",".join(str(c) for c in CONC), "x_label": "Concentration (µg/mL)"},
    ).json()

    lcms_bytes = _generate_synthetic_mzml(num_scans=6)
    lcms = client.post("/api/lcms/sessions", files={"file": ("run1.mzML", lcms_bytes, "application/octet-stream")}).json()["session_id"]
    _tag(lcms)
    client.post(f"/api/lcms/sessions/{lcms}/deconvolute", json={"rt_min": 0.1, "polarity": "positive"})
    feature = {"id": "f1", "session_id": lcms, "source_file": "run1.mzML", "mz": 500.0, "tolerance": 0.01, "polarity": "positive",
               "rt_apex": 0.1, "rt_start": 0.05, "rt_end": 0.2, "height": 1234.5, "area": 98.76, "baseline": 1.0, "n_points": 4,
               "source": "eic", "label": "M1", "created_at": "2026-09-23T00:00:00Z"}
    r = client.post("/api/automation/actions/lcms.export_feature_table_csv/execute", json={"rows": [feature]})
    assert r.status_code == 200, r.text

    resp = client.post("/api/publication/si-package", json={"experiment_tag": TAG, "include_raw_files": True})
    assert resp.status_code == 200, resp.text
    zf = zipfile.ZipFile(io.BytesIO(resp.content))
    return {"zip": zf, "peaks": peaks, "fit": fit, "mic": mic, "ftir_bytes": ftir_bytes, "plate_bytes": plate_bytes,
            "lcms_bytes": lcms_bytes, "feature": feature}


def _sheet(zf, name):
    xlsx = next(n for n in zf.namelist() if n.endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(zf.read(xlsx)))
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header_i = next(i for i, r in enumerate(rows) if r and r[0] == "File")
    header = rows[header_i]
    return [dict(zip(header, r)) for r in rows[header_i + 1:] if any(v is not None for v in r)]


def test_ftir_peaks_match_app_output(run):
    rows = _sheet(run["zip"], "FTIR_Peaks")
    assert [r["Wavenumber (cm-1)"] for r in rows] == pytest.approx([p["wn"] for p in run["peaks"]["peaks"]])


def test_ftir_fit_components_match(run):
    rows = _sheet(run["zip"], "FTIR_Fits")
    assert [r["Center (cm-1)"] for r in rows] == pytest.approx([c["center"] for c in run["fit"]["components"]])
    assert [r["Area %"] for r in rows] == pytest.approx([c["area_percent"] for c in run["fit"]["components"]])


def test_mic_fit_matches_app_output(run):
    (row,) = _sheet(run["zip"], "Plate_4PL")
    fit = run["mic"]["result"]["four_pl"]
    assert row["IC50"] == pytest.approx(fit["ic50"])
    assert row["IC50 SE"] == pytest.approx(fit["ic50_se"])
    assert row["Unit"] == "µg/mL"
    assert row["R²"] == pytest.approx(fit["r_squared"])


def test_lcms_tables_present(run):
    (feat,) = _sheet(run["zip"], "LCMS_Features")
    assert feat["Area (counts·min)"] == pytest.approx(98.76)
    assert feat["m/z"] == pytest.approx(500.0)
    assert _sheet(run["zip"], "LCMS_Deconvolution") is not None


def test_methods_describe_the_actual_settings(run):
    methods = run["zip"].read(next(n for n in run["zip"].namelist() if n.endswith("Methods.md"))).decode()
    assert "airPLS" in methods and "12345" in methods
    assert "pseudo-Voigt" in methods
    assert "4PL" in methods or "four-parameter" in methods
    # nothing from the old hard-coded text
    assert "R^2 \ge 0.98" not in methods and "10 ppm" not in methods


def test_manifest_hashes_and_raw_files(run):
    zf = run["zip"]
    manifest = json.loads(zf.read("manifest.json"))
    by_name = {s["file"]: s for s in manifest["sessions"]}
    assert by_name["pla_film.csv"]["sha256"] == hashlib.sha256(run["ftir_bytes"]).hexdigest()
    assert by_name["mic_plate.csv"]["sha256"] == hashlib.sha256(run["plate_bytes"]).hexdigest()
    assert zf.read("raw/ftir/pla_film.csv") == run["ftir_bytes"]
    assert zf.read("raw/lcms/run1.mzML") == run["lcms_bytes"]
    assert {r["kind"] for r in manifest["results"]} >= {"peaks", "fit", "mic", "deconvolution", "feature_table"}
    assert manifest["app_version"].startswith("0.1.0+")


def test_sessions_without_results_give_a_note_not_numbers():
    sid = client.post("/api/ftir/sessions", files={"file": ("bare.csv", b"\n".join(f"{600+i},{0.1}".encode() for i in range(30)), "text/csv")}).json()["session_id"]
    resp = client.post("/api/publication/si-package", json={"session_ids": [sid]})
    zf = zipfile.ZipFile(io.BytesIO(resp.content))
    xlsx = next(n for n in zf.namelist() if n.endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(zf.read(xlsx)))
    assert wb.sheetnames == ["SI_Summary"]
    text = " ".join(str(c) for row in wb["SI_Summary"].iter_rows(values_only=True) for c in row if c)
    assert "No recorded analyses" in text

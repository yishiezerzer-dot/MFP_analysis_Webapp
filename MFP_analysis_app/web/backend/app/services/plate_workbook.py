"""Calculation workbook for a plate: the raw plate plus every derived number as a live Excel formula.

Opening it in Excel (or LibreOffice) recomputes blank, growth control, mean, SD and % growth from
the raw OD values, so the analysis can be audited or redone by hand. An experiment workbook holds
several plates (sheets prefixed "P1 ", "P2 ", ...) plus a MIC reading grid of % growth values.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Sequence, Tuple

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from lab_gui.plate_gen5 import ROWS, PlateRead
from lab_gui.plate_mic import PlateLayout, well_concentration

_BOLD = Font(bold=True)


def _write_plate(wb, read: PlateRead, layout: PlateLayout, analysis: Dict[str, Any], *, subtract_blank: bool,
                 prefix: str = "") -> None:
    raw_name, ctl_name = f"{prefix}Raw plate", f"{prefix}Controls"
    excluded = set(layout.excluded)
    usable = lambda w: w not in excluded and read.values.get(w) is not None  # noqa: E731

    def raw_ref(well: str) -> str:
        row, col = well[0], int(well[1:])
        return f"'{raw_name}'!{get_column_letter(1 + col)}{3 + ROWS.index(row)}"

    raw = wb.create_sheet(raw_name)
    raw["A1"] = "OD as imported" + (f" ({read.label} nm)" if read.label else "")
    raw["A1"].font = _BOLD
    for c in range(1, 13):
        raw.cell(row=2, column=1 + c, value=c).font = _BOLD
    for i, r in enumerate(ROWS):
        raw.cell(row=3 + i, column=1, value=r).font = _BOLD
        for c in range(1, 13):
            raw.cell(row=3 + i, column=1 + c, value=read.values.get(f"{r}{c}"))
    meta_row = 12
    for key, value in (read.metadata or {}).items():
        raw.cell(row=meta_row, column=1, value=key)
        raw.cell(row=meta_row, column=2, value=value)
        meta_row += 1

    lay = wb.create_sheet(f"{prefix}Layout")
    lay.append(["Well", "Group", "Role", f"Concentration ({layout.dilution.unit})", "Excluded"])
    roles: Dict[str, tuple] = {}
    for g in layout.groups:
        for w in g.wells:
            roles[w] = (g.name, g.kind)
    for w in layout.growth_control:
        roles[w] = ("", "growth control")
    for w in layout.blank:
        roles[w] = ("", "blank")
    for r in ROWS:
        for c in range(1, 13):
            w = f"{r}{c}"
            name, role = roles.get(w, ("", ""))
            conc = well_concentration(w, layout.dilution) if role in ("sample", "reference") else None
            lay.append([w, name, role, conc, "yes" if w in excluded else ""])

    ctl = wb.create_sheet(ctl_name)
    blank_wells = [w for w in layout.blank if usable(w)]
    ctl["A1"] = "Blank (mean of blank wells)" if (subtract_blank and blank_wells) else "Blank (not subtracted)"
    ctl["B1"] = f"=AVERAGE({','.join(raw_ref(w) for w in blank_wells)})" if (subtract_blank and blank_wells) else 0
    ctl["C1"] = ", ".join(blank_wells)
    ctl.append([])
    ctl.append(["Group", "Growth control mean (blank-subtracted)", "Wells"])
    gc_row: Dict[str, int] = {}
    for g in analysis["groups"]:
        wells = g["growth_control"]["wells"]
        ctl.append([g["name"], f"=AVERAGE({','.join(raw_ref(w) for w in wells)})-'{ctl_name}'!$B$1" if wells else None,
                    ", ".join(wells)])
        gc_row[g["name"]] = ctl.max_row

    calc = wb.create_sheet(f"{prefix}Calculation")
    header = ["Group", "Concentration", "Unit", "n", "Mean OD (blank-subtracted)", "SD", "% growth", "Wells used"]
    calc.append(header)
    for cell in calc[1]:
        cell.font = _BOLD
    for g in analysis["groups"]:
        for p in g["points"]:
            used = [x["well"] for x in p["wells"] if not x["excluded"] and usable(x["well"])]
            refs = ",".join(raw_ref(w) for w in used)
            row = calc.max_row + 1
            calc.append([
                g["name"],
                p["concentration"],
                layout.dilution.unit,
                f"=COUNT({refs})" if used else 0,
                f"=AVERAGE({refs})-'{ctl_name}'!$B$1" if used else None,
                f"=STDEV.S({refs})" if len(used) > 1 else None,
                f"=E{row}/'{ctl_name}'!$B${gc_row[g['name']]}*100" if used and g["growth_control"]["wells"] else None,
                ", ".join(used),
            ])

    chk = wb.create_sheet(f"{prefix}Checks")
    chk.append(["Level", "Check"])
    for c in analysis["checks"]:
        chk.append([c["level"], c["message"]])


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_workbook(read: PlateRead, layout: PlateLayout, analysis: Dict[str, Any], *, subtract_blank: bool) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_plate(wb, read, layout, analysis, subtract_blank=subtract_blank)
    wb.move_sheet("Calculation", offset=-2)  # Raw plate, Calculation, Layout, Controls, Checks
    return _save(wb)


def build_experiment_workbook(tag: str, plates: Sequence[Tuple[str, PlateRead, PlateLayout, Dict[str, Any]]], *,
                              subtract_blank: bool) -> bytes:
    """`plates`: (display name, read, layout, analysis) per plate, in the order they are listed."""
    wb = openpyxl.Workbook()
    grid = wb.active
    grid.title = "MIC grid"
    grid.append([f"Experiment: {tag}", "% growth of each compound's own growth control (values; formulas are on the plate sheets)"])
    grid["A1"].font = _BOLD
    concs: List[float] = sorted({p["concentration"] for _, _, _, a in plates for g in a["groups"] for p in g["points"]},
                                reverse=True)
    unit = plates[0][3]["unit"] if plates else ""
    grid.append(["Compound", "Plate", "Replicate"] + [f"{c:g} {unit}" for c in concs])
    for cell in grid[2]:
        cell.font = _BOLD
    for i, (name, read, layout, analysis) in enumerate(plates, start=1):
        for g in analysis["groups"]:
            by_conc = {p["concentration"]: p for p in g["points"]}
            grid.append([g["name"], f"P{i} {name}", "mean"] + [
                (by_conc[c]["percent_growth"] if c in by_conc else None) for c in concs
            ])
            wells_by_rep: Dict[str, Dict[float, Any]] = {}
            for p in g["points"]:
                for w in p["wells"]:
                    rep = w["well"][0] if layout.dilution.direction == "columns" else w["well"][1:]
                    wells_by_rep.setdefault(rep, {})[p["concentration"]] = None if w["excluded"] else w["percent_growth"]
            for rep, vals in wells_by_rep.items():
                label = f"row {rep}" if layout.dilution.direction == "columns" else f"column {rep}"
                grid.append(["", "", label] + [vals.get(c) for c in concs])
    for i, (_, read, layout, analysis) in enumerate(plates, start=1):
        _write_plate(wb, read, layout, analysis, subtract_blank=subtract_blank, prefix=f"P{i} ")
    return _save(wb)

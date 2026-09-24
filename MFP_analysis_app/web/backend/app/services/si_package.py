"""Supplementary Information package built only from recorded analysis results.

Every number in the tables comes from `analysis_results` rows (what the app actually computed),
and the methods text is generated from the parameters stored with those rows.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font

from ..db import list_results
from ..provenance import app_version, session_input_sha256

# Kinds where only the latest result per session is meaningful (a re-run supersedes it);
# fits, integrations and deconvolutions of different regions are all kept.
_LATEST_ONLY = {"peaks", "mic", "feature_table"}
_BOLD = Font(bold=True)


def _select(results: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    by_kind: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    latest: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for r in results:  # ordered by created_at
        if r["kind"] in _LATEST_ONLY:
            latest[(r["kind"], r["session_id"])] = r
        else:
            by_kind[r["kind"]].append(r)
    for (kind, _sid), r in latest.items():
        by_kind[kind].append(r)
    return by_kind


def _unit_from_label(label: str) -> str:
    match = re.search(r"\(([^()]+)\)\s*$", label or "")
    return match.group(1) if match else ""


def _add_table(wb: openpyxl.Workbook, title: str, caption: str, header: List[str], rows: Iterable[List[Any]]) -> None:
    ws = wb.create_sheet(title=title)
    ws.append([caption])
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    ws.append([])
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(row=3, column=c).font = _BOLD
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max(len(str(cell.value or "")) for cell in col[2:]) if len(col) > 2 else 10
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)


def _tables(wb: openpyxl.Workbook, sel: Dict[str, List[Dict[str, Any]]]) -> None:
    def name(r: Dict[str, Any]) -> str:
        return r.get("input_name") or r["session_id"]

    if sel.get("feature_table"):
        rows = []
        for r in sel["feature_table"]:
            for f in r["result"]["rows"]:
                rows.append([
                    name(r), f["mz"], f.get("tolerance"), f.get("polarity"), f["rt_apex"], f.get("rt_start"),
                    f.get("rt_end"), f["height"], f["area"], f.get("baseline"), f.get("n_points"),
                    f.get("label"), f.get("expected_product"), f.get("annotation"),
                ])
        _add_table(wb, "LCMS_Features", "LC-MS integrated EIC features (as exported from the feature table)",
                   ["File", "m/z", "EIC tolerance (Da)", "Polarity", "RT apex (min)", "RT start (min)", "RT end (min)",
                    "Height (counts)", "Area (counts·min)", "Baseline", "Points", "Label", "Expected product", "Annotation"],
                   rows)
    if sel.get("deconvolution"):
        rows = []
        for r in sel["deconvolution"]:
            p = r["params"]
            rt = f"{p.get('rt_min')}–{p.get('rt_max')}" if p.get("rt_max") is not None else str(p.get("rt_min"))
            for c in r["result"]["components"]:
                rows.append([name(r), rt, c["mass"], ", ".join(str(z) for z in c.get("charges", [])),
                             c.get("total_intensity"), c.get("score"), c.get("method")])
        _add_table(wb, "LCMS_Deconvolution", "LC-MS charge-state deconvolution (neutral masses)",
                   ["File", "RT (min)", "Neutral mass (Da)", "Charges", "Total intensity", "Score", "Method"], rows)
    if sel.get("peaks"):
        rows = []
        for r in sel["peaks"]:
            top = {}
            for a in r["result"].get("assignments") or []:
                cands = a.get("candidates") or []
                if cands:
                    top[round(float(a["wn"]), 6)] = cands[0]
            for pk in r["result"]["peaks"]:
                cand = top.get(round(float(pk["wn"]), 6), {})
                rows.append([name(r), pk["wn"], pk["y"], pk["prominence"], cand.get("label"), cand.get("score")])
        _add_table(wb, "FTIR_Peaks", "FTIR picked peaks (processed spectrum, absorbance units)",
                   ["File", "Wavenumber (cm-1)", "Intensity", "Prominence", "Top library assignment", "Score"], rows)
    if sel.get("fit"):
        rows = []
        for r in sel["fit"]:
            res = r["result"]
            region = f"{res['region'][0]:g}–{res['region'][1]:g}"
            for c in res["components"]:
                rows.append([name(r), region, res["profile"], c["index"], c["center"], c["fwhm"], c["area"],
                             c["area_percent"], c.get("assignment"), res.get("r2"), res.get("converged", True)])
        _add_table(wb, "FTIR_Fits", "FTIR band fitting components",
                   ["File", "Region (cm-1)", "Profile", "Component", "Center (cm-1)", "FWHM (cm-1)", "Area",
                    "Area %", "Assignment", "R²", "Converged"], rows)
    if sel.get("integration"):
        rows = [[name(r), r["result"]["region"][0], r["result"]["region"][1], r["result"]["baseline_mode"],
                 r["result"]["area"], r["result"]["height"], r["result"]["peak_wn"], r["result"].get("fwhm")]
                for r in sel["integration"]]
        _add_table(wb, "FTIR_Integrations", "FTIR band integrations",
                   ["File", "From (cm-1)", "To (cm-1)", "Baseline", "Area", "Height", "Peak (cm-1)", "FWHM (cm-1)"], rows)
    if sel.get("mic"):
        fit_rows, mean_rows = [], []
        for r in sel["mic"]:
            res, p = r["result"], r["params"]
            fit = res.get("four_pl") or {}
            fit_rows.append([
                name(r), fit.get("ic50"), fit.get("ic50_se"), _unit_from_label(p.get("x_label", "")),
                fit.get("hill_slope"), fit.get("hill_slope_se"), fit.get("top"), fit.get("bottom"), fit.get("r_squared"),
                fit.get("ic50_in_range"), bool(p.get("subtract_blank")), res.get("four_pl_skipped_reason"),
            ])
            n = res.get("sample_n") or [None] * len(res["sample_mean"])
            for label, mean, sd, count in zip(res["x_tick_labels"], res["sample_mean"], res["sample_std"], n):
                mean_rows.append([name(r), label, mean, sd, count])
        _add_table(wb, "Plate_4PL", "Plate reader dose-response (4PL) fit on replicate means",
                   ["File", "IC50", "IC50 SE", "Unit", "Hill slope", "Hill SE", "Top", "Bottom", "R²",
                    "IC50 within tested range", "Blank subtracted", "No fit because"], fit_rows)
        _add_table(wb, "Plate_Means", "Plate reader replicate means",
                   ["File", "Concentration", "Mean", "SD", "n"], mean_rows)


def _distinct(values: Iterable[str]) -> List[str]:
    return sorted({v for v in values if v})


def _methods(sel: Dict[str, List[Dict[str, Any]]], sessions: List[Dict[str, Any]], version: str) -> str:
    import numpy
    import pandas
    import pyteomics
    import scipy

    out = ["# Supplementary methods (generated from the recorded analysis settings)", ""]
    ftir_pre = []
    for kind in ("peaks", "fit", "integration"):
        for r in sel.get(kind, []):
            p = r["params"]
            steps = []
            if p.get("mode") == "transmittance":
                steps.append("transmittance was converted to absorbance (A = −log₁₀T)")
            if int(p.get("smoothing_window") or 0) >= 3:
                steps.append(f"Savitzky–Golay smoothing (window {p['smoothing_window']}, order {p.get('poly_order')})")
            base = p.get("baseline") or "none"
            if base == "airpls":
                steps.append(f"airPLS baseline correction (λ = {p.get('baseline_lambda'):g})")
            elif base == "asls":
                steps.append(f"asymmetric least-squares baseline (λ = {p.get('baseline_lambda'):g}, p = {p.get('baseline_p')})")
            elif base == "rubberband":
                steps.append("rubber-band (lower convex hull) baseline")
            elif base == "polyfit":
                steps.append("iterative quadratic (ModPoly) baseline")
            if p.get("atr_correction"):
                steps.append(f"approximate ATR correction (intensity × ν/ν_ref, n_crystal = {p.get('atr_n_crystal')})")
            if (p.get("normalize") or "none") != "none":
                steps.append(f"{p['normalize']} normalisation")
            ftir_pre.append("; ".join(steps) if steps else "no preprocessing")
    if ftir_pre:
        out += ["## FTIR", "", "Preprocessing: " + " | ".join(_distinct(ftir_pre)) + "."]
        for r in sel.get("peaks", []):
            p = r["params"]
            how = "minima of the second derivative" if p.get("second_derivative") else "local maxima"
            line = (f"Peaks ({r['input_name']}) were picked as {how} with minimum prominence {p.get('min_prominence')} "
                    f"and minimum separation {p.get('min_distance_cm1')} cm⁻¹")
            if p.get("assign"):
                line += f"; assignments are from the built-in band library with minimum score {p.get('assign_min_score')}"
            out.append(line + ".")
        for r in sel.get("fit", []):
            p, res = r["params"], r["result"]
            profile = {"gauss": "Gaussian", "lorentz": "Lorentzian", "voigt": "pseudo-Voigt (50:50 Gaussian/Lorentzian, shared FWHM)"}[p["profile"]]
            status = "" if res.get("converged", True) else " (the fit did not converge; values are starting guesses)"
            out.append(f"Band fitting ({r['input_name']}, {p['region'][0]:g}–{p['region'][1]:g} cm⁻¹): {p['n_components']} "
                       f"{profile} components by non-linear least squares above a linear end-point baseline{status}.")
        for r in sel.get("integration", []):
            p = r["params"]
            out.append(f"Band areas ({r['input_name']}, {p['region'][0]:g}–{p['region'][1]:g} cm⁻¹) were integrated with the "
                       f"trapezoidal rule above a {p.get('baseline_mode')} baseline.")
        out.append("")
    if sel.get("mic"):
        out += ["## Plate reader", ""]
        for r in sel["mic"]:
            p, res = r["params"], r["result"]
            blank = ("blank-well means were subtracted and their standard error added in quadrature to the replicate SD; "
                     if p.get("subtract_blank") else "")
            ns = _distinct(str(n) for n in res.get("sample_n") or [])
            line = (f"{r['input_name']}: {len(p['sample_rows'])} replicate row(s) (n per concentration: {', '.join(ns) or 'n/a'}); "
                    f"{blank}")
            if res.get("four_pl"):
                line += ("a four-parameter logistic (4PL) model, y = Bottom + (Top − Bottom) / (1 + (x/IC50)^Hill), was fitted "
                         "to the replicate means at the entered concentrations by unweighted non-linear least squares; "
                         "standard errors are from the fit covariance.")
            else:
                line += f"no dose-response fit ({res.get('four_pl_skipped_reason')})."
            out.append(line)
        out.append("")
    if sel.get("feature_table") or sel.get("deconvolution"):
        out += ["## LC-MS", ""]
        if sel.get("feature_table"):
            tols = _distinct(f"{f['tolerance']:g} Da" for r in sel["feature_table"] for f in r["result"]["rows"] if f.get("tolerance"))
            out.append("Extracted-ion chromatograms summed MS1 intensity within "
                       f"±{', ±'.join(tols) if tols else 'the stated'} of each target m/z. Peaks were integrated with the trapezoidal "
                       "rule above a baseline (median of points outside the peak, or the lower edge), between the points where the "
                       "signal falls to 5 % of the apex height; the apex is the local maximum (3-point smoothed) nearest the chosen "
                       "retention time. Areas are in counts·min.")
        for r in sel.get("deconvolution", []):
            p = r["params"]
            out.append(f"Deconvolution ({r['input_name']}): charge states {p['min_charge']}–{p['max_charge']}, tolerance "
                       f"{p['tolerance']} {p['tolerance_unit']}, {p.get('polarity')} mode, peaks above "
                       f"{100 * float(p.get('min_rel_intensity', 0)):g} % of the base peak.")
        out.append("")
    out += [
        "## Software and data",
        "",
        f"MFP Analysis app {version} (results were recorded with: {', '.join(_distinct(r['app_version'] for rs in sel.values() for r in rs)) or version}); "
        f"numpy {numpy.__version__}, scipy {scipy.__version__}, pandas {pandas.__version__}, pyteomics {getattr(pyteomics, '__version__', 'n/a')}.",
        "",
        "Input files (SHA-256):",
        "",
    ]
    out += [f"- {s['display_name']}: `{s['sha256'] or 'file missing'}`" for s in sessions]
    return "\n".join(out) + "\n"


def build_si_package(
    session_records: List[Dict[str, Any]],
    *,
    experiment_tag: str,
    include_raw_files: bool,
    figure_pdf: Optional[bytes] = None,
) -> bytes:
    version = app_version()
    now = datetime.now(timezone.utc)
    sessions = []
    for rec in session_records:
        sessions.append({
            "session_id": rec["session_id"],
            "module": rec["module"],
            "file": rec["display_name"],
            "display_name": rec["display_name"],
            "sha256": session_input_sha256(rec),
            "uploaded_at": rec.get("created_at"),
            "path": rec["file_path"],
        })
    results = list_results(session_ids=[s["session_id"] for s in sessions])
    sel = _select(results)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SI_Summary"
    ws.append(["Supplementary information — data index"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([f"Experiment: {experiment_tag or '(selected sessions)'}"])
    ws.append([f"Generated: {now.strftime('%Y-%m-%d %H:%M UTC')} with MFP Analysis {version}"])
    if not results:
        ws.append(["No recorded analyses for these sessions: run the analyses in the app (or export the LCMS feature "
                   "table) and build the package again."])
    ws.append([])
    ws.append(["File", "Module", "SHA-256", "Uploaded", "Session ID"])
    for c in range(1, 6):
        ws.cell(row=ws.max_row, column=c).font = _BOLD
    for s in sessions:
        ws.append([s["file"], s["module"], s["sha256"], s["uploaded_at"], s["session_id"]])
    _tables(wb, sel)
    xlsx = io.BytesIO()
    wb.save(xlsx)

    manifest = {
        "generated_at": now.isoformat(),
        "app_version": version,
        "experiment_tag": experiment_tag,
        "sessions": [{k: v for k, v in s.items() if k not in ("path", "display_name")} for s in sessions],
        "results": [
            {k: r[k] for k in ("id", "session_id", "module", "kind", "params", "result", "input_name", "input_sha256",
                               "app_version", "created_at")}
            for rs in sel.values() for r in rs
        ],
    }
    stem = re.sub(r"[^\w\-.]", "_", experiment_tag) if experiment_tag else "SI"
    readme = [
        f"Supplementary information package: {experiment_tag or 'selected sessions'}",
        f"Generated {now.strftime('%Y-%m-%d %H:%M UTC')} by MFP Analysis {version}.",
        "",
        f"{stem}_Data_Tables.xlsx  tables built only from analyses recorded in the app",
        f"{stem}_Methods.md        methods text generated from the recorded settings",
        "manifest.json             input file SHA-256 hashes and every result with its parameters",
    ]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{stem}_Data_Tables.xlsx", xlsx.getvalue())
        zf.writestr(f"{stem}_Methods.md", _methods(sel, sessions, version).encode("utf-8"))
        zf.writestr("manifest.json", json.dumps(manifest, indent=2, default=str))
        if figure_pdf:
            zf.writestr(f"figures/{stem}_Figure.pdf", figure_pdf)
            readme.append("figures/                  assembled figure (PDF; panels embedded as images)")
        if include_raw_files:
            missing = []
            used: set = set()
            for s in sessions:
                path = Path(s["path"])
                if path.is_file():
                    arcname = f"raw/{s['module']}/{s['file']}"
                    if arcname in used:  # two sessions uploaded under the same file name
                        arcname = f"raw/{s['module']}/{s['session_id'][:8]}_{s['file']}"
                    used.add(arcname)
                    zf.write(path, arcname)
                else:
                    missing.append(s["file"])
            readme.append("raw/                      the original input files")
            if missing:
                readme.append(f"  (not included, file no longer on the server: {', '.join(missing)})")
        zf.writestr("README.txt", "\n".join(readme) + "\n")
    return buf.getvalue()
